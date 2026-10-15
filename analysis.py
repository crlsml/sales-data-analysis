
import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import matplotlib.pyplot as plt

with open("data/ventas.csv", "r", encoding="utf-8") as archivo:
    lector = csv.DictReader(archivo)
    ventas = list(lector)

print("REPORTE DE VENTAS")
print("-----------------")
print("Total de registros:", len(ventas))

print("\nPrimer registro:")
print(ventas[0])

print("\nINGRESOS POR VENTA")
print("------------------")

for venta in ventas:
    unidades = int(venta["unidades"])
    precio = float(venta["precio_unitario"])
    ingresos = unidades * precio
    print(venta["producto"], ":", f"${ingresos:.2f}")


print("\nRESUMEN GENERAL")
print("----------------")

ingreso_total = 0

for venta in ventas:
    unidades = int(venta["unidades"])
    precio = float(venta["precio_unitario"])
    ingresos = unidades * precio
    ingreso_total += ingresos

print(f"Ingreso total: ${ingreso_total:,.2f}")


print("\nINGRESOS POR PRODUCTO")
print("----------------------")

ingresos_por_producto = {}

for venta in ventas:
    producto = venta["producto"]
    unidades = int(venta["unidades"])
    precio = float(venta["precio_unitario"])

    ingresos = unidades * precio

    if producto not in ingresos_por_producto:
        ingresos_por_producto[producto] = 0

    ingresos_por_producto[producto] += ingresos

for producto, ingresos in ingresos_por_producto.items():
    print(f"{producto}: ${ingresos:,.2f}")



print("\nRANKING DE PRODUCTOS")
print("--------------------")

productos_ordenados = sorted(
    ingresos_por_producto.items(),
    key=lambda elemento: elemento[1],
    reverse=True
)

for posicion, (producto, ingresos) in enumerate(productos_ordenados, start=1):
    print(f"{posicion}. {producto}: ${ingresos:,.2f}")



print("\nANALISIS POR VENDEDOR")
print("---------------------")

ventas_por_vendedor = {}

for venta in ventas:
    vendedor = venta["vendedor"]
    unidades = int(venta["unidades"])
    precio = float(venta["precio_unitario"])

    ingresos = unidades * precio

    if vendedor not in ventas_por_vendedor:
        ventas_por_vendedor[vendedor] = {
            "ingresos": 0,
            "operaciones": 0,
            "unidades": 0
        }

    ventas_por_vendedor[vendedor]["ingresos"] += ingresos
    ventas_por_vendedor[vendedor]["operaciones"] += 1
    ventas_por_vendedor[vendedor]["unidades"] += unidades

for vendedor, datos in ventas_por_vendedor.items():
    print(f"\nVendedor: {vendedor}")
    print(f"Ingresos: ${datos['ingresos']:,.2f}")
    print(f"Operaciones: {datos['operaciones']}")
    print(f"Unidades vendidas: {datos['unidades']}")

mejor_vendedor = max(
    ventas_por_vendedor,
    key=lambda vendedor: ventas_por_vendedor[vendedor]["ingresos"]
)

datos_mejor_vendedor = ventas_por_vendedor[mejor_vendedor]

print("\nVENDEDOR CON MAYORES INGRESOS")
print("-----------------------------")
print(f"Vendedor: {mejor_vendedor}")
print(f"Ingresos: ${datos_mejor_vendedor['ingresos']:,.2f}")



print("\nANALISIS MENSUAL")
print("----------------")

ingresos_por_mes = {}

for venta in ventas:
    fecha = venta["fecha"]
    mes = fecha[:7]

    unidades = int(venta["unidades"])
    precio = float(venta["precio_unitario"])
    ingresos = unidades * precio

    if mes not in ingresos_por_mes:
        ingresos_por_mes[mes] = 0

    ingresos_por_mes[mes] += ingresos

for mes, ingresos in sorted(ingresos_por_mes.items()):
    print(f"{mes}: ${ingresos:,.2f}")


mejor_mes = max(
    ingresos_por_mes,
    key=ingresos_por_mes.get
)

print("\nMES CON MAYORES INGRESOS")
print("------------------------")
print(f"Mes: {mejor_mes}")
print(f"Ingresos: ${ingresos_por_mes[mejor_mes]:,.2f}")



print("\n" + "=" * 40)
print("       RESUMEN EJECUTIVO")
print("=" * 40)

total_operaciones = len(ventas)

total_unidades = sum(
    int(venta["unidades"]) for venta in ventas
)

producto_top = max(
    ingresos_por_producto,
    key=ingresos_por_producto.get
)

mes_top = max(
    ingresos_por_mes,
    key=ingresos_por_mes.get
)

print(f"Ingresos totales: ${ingreso_total:,.2f}")
print(f"Operaciones registradas: {total_operaciones}")
print(f"Unidades vendidas: {total_unidades}")

print(
    f"Producto destacado: {producto_top} "
    f"(${ingresos_por_producto[producto_top]:,.2f})"
)

print(
    f"Mes destacado: {mes_top} "
    f"(${ingresos_por_mes[mes_top]:,.2f})"
)

print("=" * 40)


import os

os.makedirs("reports", exist_ok=True)

# Exportar ingresos por producto
with open(
    "reports/resumen_productos.csv",
    "w",
    newline="",
    encoding="utf-8-sig"
) as archivo:
    escritor = csv.writer(archivo)
    escritor.writerow(["Producto", "Ingresos"])

    for producto, ingresos in sorted(
        ingresos_por_producto.items(),
        key=lambda elemento: elemento[1],
        reverse=True
    ):
        escritor.writerow([producto, ingresos])

print("\nArchivo CSV exportado: reports/resumen_productos.csv")



# Crear el libro de Excel
libro = Workbook()

# Hoja 1: Resumen ejecutivo
hoja_resumen = libro.active
hoja_resumen.title = "Resumen"

hoja_resumen.append(["Indicador", "Resultado"])
hoja_resumen.append(["Ingresos totales", ingreso_total])
hoja_resumen.append(["Operaciones registradas", len(ventas)])
hoja_resumen.append([
    "Unidades vendidas",
    sum(int(venta["unidades"]) for venta in ventas)
])

producto_top = max(
    ingresos_por_producto,
    key=ingresos_por_producto.get
)

hoja_resumen.append(["Producto con mayores ingresos", producto_top])
hoja_resumen.append([
    "Ingresos del producto destacado",
    ingresos_por_producto[producto_top]
])

mes_top = max(ingresos_por_mes, key=ingresos_por_mes.get)

hoja_resumen.append(["Mes con mayores ingresos", mes_top])
hoja_resumen.append(["Ingresos del mejor mes", ingresos_por_mes[mes_top]])


# Hoja 2: Productos
hoja_productos = libro.create_sheet("Productos")
hoja_productos.append(["Producto", "Ingresos"])

for producto, ingresos in sorted(
    ingresos_por_producto.items(),
    key=lambda elemento: elemento[1],
    reverse=True
):
    hoja_productos.append([producto, ingresos])


# Hoja 3: Vendedores
hoja_vendedores = libro.create_sheet("Vendedores")
hoja_vendedores.append([
    "Vendedor",
    "Ingresos",
    "Operaciones",
    "Unidades vendidas"
])

for vendedor, datos in sorted(
    ventas_por_vendedor.items(),
    key=lambda elemento: elemento[1]["ingresos"],
    reverse=True
):
    hoja_vendedores.append([
        vendedor,
        datos["ingresos"],
        datos["operaciones"],
        datos["unidades"]
    ])


# Hoja 4: Mensual
hoja_mensual = libro.create_sheet("Mensual")
hoja_mensual.append(["Mes", "Ingresos"])

for mes, ingresos in sorted(ingresos_por_mes.items()):
    hoja_mensual.append([mes, ingresos])


# Dar formato a las hojas
for hoja in libro.worksheets:
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="217346"
        )

    for columna in hoja.columns:
        ancho = max(
            len(str(celda.value or "")) for celda in columna
        )
        hoja.column_dimensions[columna[0].column_letter].width = ancho + 3


# Formatear los importes monetarios
for fila in range(2, hoja_resumen.max_row + 1):
    if hoja_resumen.cell(fila, 1).value in [
        "Ingresos totales",
        "Ingresos del producto destacado",
        "Ingresos del mejor mes"
    ]:
        hoja_resumen.cell(fila, 2).number_format = '$#,##0.00'

for hoja in [hoja_productos, hoja_vendedores, hoja_mensual]:
    for fila in range(2, hoja.max_row + 1):
        hoja.cell(fila, 2).number_format = '$#,##0.00'


# Guardar el archivo
libro.save("reports/reporte_ventas.xlsx")

print("\nExcel exportado correctamente:")
print("reports/reporte_ventas.xlsx")


# Crear gráfica de ingresos por producto
os.makedirs("reports", exist_ok=True)

productos = list(ingresos_por_producto.keys())
ingresos = list(ingresos_por_producto.values())

plt.figure(figsize=(10, 6))
plt.bar(productos, ingresos)
plt.title("Ingresos por producto")
plt.xlabel("Producto")
plt.ylabel("Ingresos ($)")
plt.xticks(rotation=25, ha="right")
plt.tight_layout()

plt.savefig("reports/ingresos_por_producto.png", dpi=300)
plt.close()

print("Gráfica guardada en reports/ingresos_por_producto.png")


# Crear gráfica de ingresos mensuales
meses = list(ingresos_por_mes.keys())
ingresos_mensuales = list(ingresos_por_mes.values())

plt.figure(figsize=(8, 5))
plt.plot(meses, ingresos_mensuales, marker="o")
plt.title("Evolución mensual de ingresos")
plt.xlabel("Mes")
plt.ylabel("Ingresos ($)")
plt.grid(True, alpha=0.3)
plt.tight_layout()

plt.savefig("reports/ingresos_mensuales.png", dpi=300)
plt.close()

print("Gráfica guardada en reports/ingresos_mensuales.png")


# Crear gráfica de ingresos por vendedor
vendedores = list(ventas_por_vendedor.keys())

ingresos_vendedores = [
    datos["ingresos"]
    for datos in ventas_por_vendedor.values()
]

plt.figure(figsize=(8, 5))
plt.bar(vendedores, ingresos_vendedores)
plt.title("Ingresos por vendedor")
plt.xlabel("Vendedor")
plt.ylabel("Ingresos ($)")
plt.tight_layout()

plt.savefig("reports/ingresos_por_vendedor.png", dpi=300)
plt.close()

print("Gráfica guardada en reports/ingresos_por_vendedor.png")
